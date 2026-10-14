"""Acquire immutable source archives, retaining metadata and verifying counts."""

import csv
from datetime import datetime, timezone
import hashlib
from http.client import HTTPException
import io
import json
import math
from pathlib import Path
import re
import time
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import zipfile

BASE = 'https://data.un.org/legacy/'
TABLES = {101: 24, 102: 24, 201: 16, 202: 16, 204: 26, 205: 26}
FX_URL = 'https://unstats.un.org/unsd/amaapi/api/file/30'
WORLD_URL = 'https://unstats.un.org/unsd/amaapi/api/file/7'
REQUIRED = {'Country or Area Code', 'Country or Area', 'Fiscal Year', 'Series',
            'Currency', 'SNA System', 'Fiscal Year Type', 'Value', 'Item Code', 'SNA93 Item Code'}
IDENTITY = ('Group Code', 'Country or Area Code', 'Fiscal Year', 'Series',
            'Currency', 'SNA System', 'Fiscal Year Type', 'Base Year')


def read_un_csv(blob: bytes) -> list[dict]:
    """Read the UN's ZIP CSV, including its duplicate footnote columns and footer."""
    if not zipfile.is_zipfile(io.BytesIO(blob)):
        raise ValueError('Expected a ZIP archive, not HTML or an error response')
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        names = [n for n in archive.namelist() if n.lower().endswith('.csv')]
        if len(names) != 1:
            raise ValueError('ZIP must contain exactly one CSV')
        text = archive.read(names[0]).decode('utf-8-sig')
    reader = csv.reader(io.StringIO(text, newline=''))
    header = ['Fiscal Year' if name == 'Year' else name for name in next(reader, [])]
    missing = REQUIRED - set(header)
    if missing:
        raise ValueError('Missing required columns: ' + ', '.join(sorted(missing)))
    records, footnotes = [], {}
    for line in reader:
        if not line or not any(line):
            continue
        if len(line) <= 2:
            if len(line) == 2 and line[0].strip().isdigit():
                footnotes[line[0].strip()] = line[1].strip()
            elif line[0].strip().lower().replace('_', '') not in ('footnoteseqid', 'footnotes', 'footnote'):
                raise ValueError('Unexpected short CSV row: ' + repr(line))
            continue
        if len(line) != len(header):
            raise ValueError('Unexpected CSV row width')
        row = {}
        for name, value in zip(header, line):
            value = value.strip()
            row[name] = '; '.join(filter(None, (row.get(name, ''), value)))
        if not row['Fiscal Year'].isdigit():
            raise ValueError('Invalid observation year')
        records.append(row)
    for row in records:
        references = re.findall(r'\d+', row.get('Value Footnotes', ''))
        row['Footnote text'] = '; '.join(footnotes.get(n, 'Unresolved footnote ' + n)
                                        for n in dict.fromkeys(references))
    return records


def validate_rows(rows: list[dict], expected_count: int) -> None:
    if len(rows) != expected_count or not 0 < expected_count < 100000:
        raise ValueError(f'Record count mismatch/truncation: {len(rows)} vs {expected_count}')
    seen = set()
    for row in rows:
        expected_item = str(TABLES.get(int(row.get('Group Code', '0')), 24))
        if (row['Item Code'] != expected_item or row['SNA93 Item Code'] != 'B.1*g'
                or 'gross domestic product' not in row['Item'].casefold()):
            raise ValueError('Export includes an item other than GDP')
        key = tuple(row.get(k, '') for k in IDENTITY)
        if key in seen:
            raise ValueError('duplicate observation: ' + repr(key))
        seen.add(key)


def read_exchange_rates(blob: bytes, year: int = 2020) -> dict:
    """Read the IMF-based measure only; retain country, currency and source note."""
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    rates = {}
    rows = iter(book['Download-XPop'].values)
    next(rows); next(rows)
    header = next(rows)
    column = header.index(year)
    for row in rows:
        if row[3] != 'IMF based exchange rate' or row[column] is None:
            continue
        rate = float(row[column])
        if not math.isfinite(rate) or rate <= 0:
            continue
        code = str(int(row[0]))
        if code in rates:
            raise ValueError('duplicate exchange rate: ' + code)
        rates[code] = {'country': row[1], 'currency': row[2], 'rate': rate,
                       'year': year, 'measure': row[3], 'source_note': ''}
    note_rows = iter(book['Download-XPopNotes'].values)
    next(note_rows); next(note_rows)
    note_header = next(note_rows)
    note_column = note_header.index(year)
    for row in note_rows:
        code = str(int(row[0])) if isinstance(row[0], (int, float)) else ''
        if code in rates and row[3] == 'IMF based exchange rate' and row[note_column]:
            prior = rates[code]['source_note']
            rates[code]['source_note'] = '; '.join(filter(None, (prior, str(row[4]))))
    book.close()
    return rates


def read_world_reference(blob: bytes) -> dict:
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    rows = iter(book.active.values)
    next(rows); next(rows)
    header = next(rows)
    matches = [row for row in rows if row[1] == 'World' and row[2] == 'Gross Domestic Product (GDP)']
    book.close()
    if len(matches) != 1:
        raise ValueError('Expected exactly one reference world GDP row')
    values = {int(year): float(value) for year, value in zip(header[3:], matches[0][3:])
              if isinstance(year, int) and isinstance(value, (int, float)) and math.isfinite(value) and value > 0}
    if not values:
        raise ValueError('Empty reference world GDP')
    first, last = min(values), max(values)
    points = [{'year': y, 'value': values.get(y),
               'index': 100 * values[y] / values[first] if y in values else None,
               'growth': 100 * (values[y] / values[y - 1] - 1) if y in values and y - 1 in values else None,
               'notes': '国連AMA。推計・補完・過去への接続を含む。2008SNA公式値だけの集計ではない。'}
              for y in range(first, last + 1)]
    return {'id': 'un-world-reference', 'kind': 'reference', 'label': '参考：国連世界合計（推計・補完を含む）',
            'direct_2008_only': False, 'first': first, 'last': last, 'index_base': first,
            'currency': '2020年価格米ドル', 'base_year': '2020', 'points': points,
            'cumulative_growth': 100 * (values[last] / values[first] - 1),
            'cagr': 100 * ((values[last] / values[first]) ** (1 / (last - first)) - 1)}


def request_bytes(url: str) -> bytes:
    for attempt in range(3):
        try:
            request = Request(url, headers={'User-Agent': 'world-gdp-2008sna/1.0 (research; source attribution retained)'})
            with urlopen(request, timeout=60) as response:
                return response.read()
        except (OSError, TimeoutError, HTTPException):
            if attempt == 2:
                raise
            time.sleep(2 ** attempt)
    raise AssertionError('unreachable')


def file_record(path: Path, url: str, **metadata) -> dict:
    return {'file': path.name, 'url': url, 'bytes': path.stat().st_size,
            'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), **metadata}


def fetch_snapshot(root: Path) -> Path:
    stamp = datetime.now(timezone.utc)
    snapshot = root / 'data' / 'raw' / stamp.strftime('%Y%m%dT%H%M%S%fZ')
    snapshot.mkdir(parents=True, exist_ok=False)
    files = []
    for group, item in TABLES.items():
        selection = f'group_code:{group};item_code:{item}'
        page_url = BASE + 'Data.aspx?' + urlencode({'d': 'SNA', 'f': selection})
        page = request_bytes(page_url)
        html = page.decode('utf-8-sig')
        count = re.search(r'id="spanRecordCountB">([\d,]+)</span>', html)
        view = html.split('id="divView"', 1)[-1].split('id="divSort"', 1)[0]
        columns = re.findall(r'id="v_[^"]+"', view)
        if not count or not 10 <= len(columns) <= 20:
            raise ValueError('Cannot verify source count and metadata columns')
        expected = int(count.group(1).replace(',', ''))
        page_path = snapshot / f'table-{group}-selection.html'
        page_path.write_bytes(page)
        update = re.search(r'Last update in UNdata:\s*([^<\r\n]+)', html)
        files.append(file_record(page_path, page_url, expected_rows=expected,
                                 source_update=update.group(1).strip() if update else None))
        url = BASE + 'Handlers/DownloadHandler.ashx?' + urlencode({
            'DataFilter': selection, 'DataMartId': 'SNA', 'Format': 'csv',
            'c': ','.join(str(i) for i in range(len(columns)))})
        blob = request_bytes(url)
        rows = read_un_csv(blob)
        validate_rows(rows, expected)
        archive_path = snapshot / f'table-{group}-gdp.zip'
        archive_path.write_bytes(blob)
        files.append(file_record(archive_path, url, group=group,
                                 verified_rows=len(rows), expected_rows=expected))
        print(f'Table {group}: {len(rows):,} rows; source count matched', flush=True)
    exchange = request_bytes(FX_URL)
    rates = read_exchange_rates(exchange)
    exchange_path = snapshot / 'exchange-rates.xlsx'
    exchange_path.write_bytes(exchange)
    files.append(file_record(exchange_path, FX_URL, role='supporting_exchange_rates',
                             benchmark_year=2020, verified_countries=len(rates),
                             measure='IMF based exchange rate'))
    print(f'FX: {len(rates)} countries with 2020 IMF-based rate', flush=True)
    reference_blob = request_bytes(WORLD_URL)
    reference = read_world_reference(reference_blob)
    reference_path = snapshot / 'world-reference.xlsx'
    reference_path.write_bytes(reference_blob)
    files.append(file_record(reference_path, WORLD_URL, role='reference_world_estimates',
                             direct_2008_only=False, first=reference['first'], last=reference['last']))
    manifest = {'retrieved_at': stamp.isoformat(), 'database': 'UN National Accounts Official Country Data',
                'files': files, 'notes': 'Official series; not the supplemented AMA database.'}
    (snapshot / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (snapshot.parent / 'LATEST').write_text(snapshot.name + '\n', encoding='utf-8')
    return snapshot


def load_snapshot(root: Path, snapshot: Path | None = None) -> tuple[dict, list[dict], list[dict]]:
    if snapshot is None:
        raw = root / 'data' / 'raw'
        snapshot = raw / (raw / 'LATEST').read_text(encoding='utf-8').strip()
    manifest = json.loads((snapshot / 'manifest.json').read_text(encoding='utf-8'))
    tables = {}
    for entry in manifest['files']:
        content = (snapshot / entry['file']).read_bytes()
        if hashlib.sha256(content).hexdigest() != entry['sha256']:
            raise ValueError('Source checksum mismatch: ' + entry['file'])
        if 'group' in entry:
            rows = read_un_csv(content)
            validate_rows(rows, entry['expected_rows'])
            tables[entry['group']] = rows
    return manifest, sum((tables.get(g, []) for g in (101, 201, 204)), []), sum((tables.get(g, []) for g in (102, 202, 205)), [])
