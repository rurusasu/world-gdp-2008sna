import csv
import io
import unittest
from unittest.mock import patch
from http.client import IncompleteRead
import zipfile

from gdp_report.acquire import read_un_csv, validate_rows, read_exchange_rates, read_world_reference, request_bytes
from openpyxl import Workbook


HEADER = ['Group Code', 'Country or Area Code', 'Country or Area',
          'SNA93 Table Code', 'Sub Group', 'Item Code', 'Item',
          'SNA93 Item Code', 'Fiscal Year', 'Series', 'Currency',
          'SNA System', 'Fiscal Year Type', 'Base Year', 'Value',
          'Value Footnotes', 'Value Footnotes']
ROW = ['102', '392', 'Japan', '1.2', 'Expenditures', '24',
       'Equals: GROSS DOMESTIC PRODUCT', 'B.1*g', '2020', '1000',
       'Yen', '2008', 'Western calendar year', '2020', '100', '1', '']


def archive(rows):
    stream = io.StringIO(newline='')
    csv.writer(stream).writerows(rows)
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as z:
        z.writestr('UNdata.csv', stream.getvalue().encode('utf-8-sig'))
    return out.getvalue()


class AcquisitionTests(unittest.TestCase):
    def test_interrupted_http_body_is_retried_without_accepting_partial_bytes(self):
        class Interrupted(io.BytesIO):
            def read(self, *args):
                raise IncompleteRead(b'partial', 1)
        with patch('gdp_report.acquire.urlopen', side_effect=[Interrupted(), io.BytesIO(b'complete')]), patch('gdp_report.acquire.time.sleep'):
            self.assertEqual(request_bytes('https://example.test/data'), b'complete')

    def test_world_reference_selects_world_gdp_and_never_calls_it_direct(self):
        book = Workbook()
        sheet = book.active
        sheet.append(['Title']); sheet.append([])
        sheet.append(['Region/SubregionID', 'Region/Subregion', 'IndicatorName', 1970, 1971])
        sheet.append([2, 'Africa', 'Gross Domestic Product (GDP)', 50, 60])
        sheet.append([1, 'World', 'Gross Domestic Product (GDP)', 100, 110])
        out = io.BytesIO(); book.save(out)
        reference = read_world_reference(out.getvalue())
        self.assertEqual([p['value'] for p in reference['points']], [100, 110])
        self.assertAlmostEqual(reference['points'][1]['growth'], 10)
        self.assertEqual(reference['kind'], 'reference')
        self.assertFalse(reference['direct_2008_only'])

    def test_fx_uses_only_imf_measure_and_preserves_currency_and_year(self):
        book = Workbook()
        sheet = book.active
        sheet.title = 'Download-XPop'
        sheet.append(['Exchange rates'])
        sheet.append([])
        sheet.append(['CountryID', 'Country', 'Currency', 'Measure', 2020])
        sheet.append([392, 'Japan', 'Yen', 'AMA exchange rate', 999])
        sheet.append([392, 'Japan', 'Yen', 'IMF based exchange rate', 106.77])
        notes = book.create_sheet('Download-XPopNotes')
        notes.append(['Notes'])
        notes.append([])
        notes.append(['CountryID', 'Country', 'Currency', 'Name', 'Note', 2020])
        notes.append([392, 'Japan', 'Yen', 'IMF based exchange rate', 'IMF Principal Rate', 'X'])
        out = io.BytesIO()
        book.save(out)
        rates = read_exchange_rates(out.getvalue(), year=2020)
        self.assertEqual(rates['392']['rate'], 106.77)
        self.assertEqual(rates['392']['currency'], 'Yen')
        self.assertEqual(rates['392']['source_note'], 'IMF Principal Rate')
        self.assertEqual(rates['392']['year'], 2020)

    def test_zip_preserves_country_code_and_resolves_footnotes(self):
        rows = read_un_csv(archive([HEADER, ROW, [], ['footnote_SeqID', 'Footnote'],
                                   ['1', 'Provisional estimate.']]))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['Country or Area Code'], '392')
        self.assertEqual(rows[0]['Value Footnotes'], '1')
        self.assertEqual(rows[0]['Footnote text'], 'Provisional estimate.')

    def test_html_error_is_never_accepted_as_data(self):
        with self.assertRaisesRegex(ValueError, 'ZIP'):
            read_un_csv(b'<html>Service unavailable</html>')

    def test_missing_required_metadata_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'SNA System'):
            read_un_csv(archive([['Country or Area', 'Value'], ['Japan', '100']]))

    def test_truncated_export_is_rejected(self):
        rows = read_un_csv(archive([HEADER, ROW]))
        with self.assertRaisesRegex(ValueError, 'count'):
            validate_rows(rows, expected_count=2)

    def test_duplicate_observation_is_rejected(self):
        rows = read_un_csv(archive([HEADER, ROW, ROW]))
        with self.assertRaisesRegex(ValueError, 'duplicate'):
            validate_rows(rows, expected_count=2)

    def test_non_gdp_download_is_rejected(self):
        wrong = ROW.copy()
        wrong[5] = '7'
        rows = read_un_csv(archive([HEADER, wrong]))
        with self.assertRaisesRegex(ValueError, 'GDP'):
            validate_rows(rows, expected_count=1)

    def test_nominal_table_year_heading_is_normalized(self):
        header = ['Year' if name == 'Fiscal Year' else name for name in HEADER]
        rows = read_un_csv(archive([header, ROW]))
        self.assertEqual(rows[0]['Fiscal Year'], '2020')

    def test_production_table_gdp_item_is_accepted(self):
        production = ROW.copy()
        production[0], production[3], production[5] = '205', '2.5', '26'
        rows = read_un_csv(archive([HEADER, production]))
        validate_rows(rows, expected_count=1)


if __name__ == '__main__':
    unittest.main()
